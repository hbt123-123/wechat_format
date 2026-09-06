#!/usr/bin/env python3
# -*- coding: utf-8 -*-
"""
微信聊天记录分析工具（单文件版）

输入一个 chatlog API URL，自动完成：数据下载 → 格式化 → 全量统计 → 生成心理测评任务包。

用法：
    python main.py <URL或本地JSON文件路径> [-o 输出目录]

示例：
    python main.py "http://127.0.0.1:5030/api/v1/chatlog?limit=100000"

输出（默认在 ./output 下）：
    chat_format.json       格式化后的聊天记录
    chat_analysis.xlsx     全量统计结果（语音电话、常用词、消息趋势等，每张工作表附带对应图表）
    assessment_tasks.json  心理测评任务包（供子智能体执行，包内含完整提示词与聊天记录）
    run_summary.json       运行摘要（各阶段状态与统计概要，供主智能体监测统计过程）

智能体协作说明：
    本脚本作为"统计工具"由主智能体调用并监测；心理测评由主智能体派发的子智能体
    读取 assessment_tasks.json 中的任务包完成（不再由脚本直接调用 LLM API）。
"""

import argparse
import calendar
import json
import os
import re
import sys
from collections import Counter
from datetime import datetime

import jieba
import pandas as pd
import requests
from openpyxl.chart import BarChart, LineChart, Reference
from openpyxl.chart.label import DataLabelList
from openpyxl.chart.marker import Marker

jieba.setLogLevel(30)

TOTAL_STAGES = 4


# ============================================================
# 一、运行监测器（供主智能体监测统计过程）
# ============================================================

class RunMonitor:
    """分阶段记录统计过程，最终生成 run_summary.json"""

    def __init__(self, source, output_dir):
        self.source = source
        self.output_dir = output_dir
        self.stages = []
        self.started_at = datetime.now()

    def _print(self, stage):
        status_text = {"running": "进行中", "done": "完成", "failed": "失败"}[stage["status"]]
        line = "[阶段 {}/{}] {} - {}".format(stage["stage"], TOTAL_STAGES, stage["name"], status_text)
        if stage["detail"]:
            line += "：{}".format(stage["detail"])
        print(line, flush=True)

    def report(self, name, status, detail=""):
        now = datetime.now().strftime("%Y-%m-%d %H:%M:%S")
        # 若该阶段已有"进行中"记录，则更新为最终状态
        for stage in self.stages:
            if stage["name"] == name and stage["status"] == "running":
                stage["status"] = status
                stage["detail"] = detail or stage["detail"]
                stage["finished_at"] = now
                self._print(stage)
                return
        stage = {
            "stage": len(self.stages) + 1,
            "name": name,
            "status": status,
            "detail": detail,
            "started_at": now,
            "finished_at": now,
        }
        self.stages.append(stage)
        self._print(stage)

    def fail(self, detail):
        """将当前"进行中"的阶段标记为失败"""
        for stage in reversed(self.stages):
            if stage["status"] == "running":
                stage["status"] = "failed"
                stage["detail"] = detail
                stage["finished_at"] = datetime.now().strftime("%Y-%m-%d %H:%M:%S")
                self._print(stage)
                return
        self.report("执行", "failed", detail)

    def summary(self, status, outputs, stats=None):
        result = {
            "status": status,
            "source": self.source,
            "started_at": self.started_at.strftime("%Y-%m-%d %H:%M:%S"),
            "finished_at": datetime.now().strftime("%Y-%m-%d %H:%M:%S"),
            "stages": self.stages,
            "outputs": outputs,
            "statistics_summary": stats or {},
        }
        if status == "success":
            result["next_step"] = (
                "统计流程已完成。请主智能体派发心理测评子智能体，读取 "
                "assessment_tasks.json，逐个执行 tasks 中的测评任务并汇总报告。"
            )
        return result


# ============================================================
# 二、数据获取与格式化
# ============================================================

def load_chat_messages(source, timeout):
    """从 URL（chatlog API）或本地 JSON 文件加载原始聊天记录"""
    if source.lower().startswith(("http://", "https://")):
        print("正在从 URL 下载数据：{}".format(source))
        try:
            resp = requests.get(source, timeout=timeout)
            resp.raise_for_status()
        except requests.RequestException as e:
            raise RuntimeError("下载数据失败：{}".format(e))
        try:
            data = resp.json()
        except ValueError:
            raise RuntimeError("URL 返回的内容不是有效 JSON")
    else:
        if not os.path.exists(source):
            raise RuntimeError("文件不存在：{}".format(source))
        with open(source, "r", encoding="utf-8") as f:
            try:
                data = json.load(f)
            except ValueError as e:
                raise RuntimeError("JSON 解析失败：{}".format(e))

    # 兼容裸数组 / {"data": [...]} / {"messages": [...]} 等返回结构
    if isinstance(data, list):
        messages = data
    elif isinstance(data, dict):
        for key in ("data", "messages", "list", "items"):
            if isinstance(data.get(key), list):
                messages = data[key]
                break
        else:
            raise RuntimeError("返回的 JSON 中未找到聊天记录数组")
    else:
        raise RuntimeError("返回的 JSON 格式无法识别")

    if not messages:
        raise RuntimeError("聊天记录为空，请检查 URL 参数（如 limit/time/talker）")
    return messages


def _pick(msg, *keys, default=None):
    """字段别名归一（兼容不同版本 chatlog 输出的驼峰/下划线字段）"""
    for k in keys:
        if msg.get(k) is not None:
            return msg[k]
    return default


def format_chat_messages(raw_messages):
    """原始 chatlog 消息 → [{time, sender, message}] 列表"""
    optimized_messages = []

    for msg in raw_messages:
        # 时间：ISO → YYYY-MM-DD HH:MM:SS
        time_str = _pick(msg, "time", "createTime", default="") or ""
        formatted_time = time_str
        if time_str:
            try:
                dt = datetime.fromisoformat(str(time_str).replace("Z", "+00:00"))
                formatted_time = dt.strftime("%Y-%m-%d %H:%M:%S")
            except ValueError:
                formatted_time = time_str

        # 发送者
        if _pick(msg, "isSelf", "is_self", default=False):
            formatted_sender = "自己"
        else:
            formatted_sender = _pick(msg, "senderName", "sender_name", "talker", default="") or "对方"

        # 消息内容（按类型映射）
        msg_type = _pick(msg, "type", "msg_type", default=0)
        if msg_type == 1:
            msg_content = msg.get("content", "") or ""
        elif msg_type == 3:
            msg_content = "[photo]"
        elif msg_type == 34:
            msg_content = "[voice]"
        elif msg_type == 43:
            msg_content = "[video]"
        elif msg_type == 47:
            msg_content = "[emoji]"
        elif msg_type == 49:
            contents = msg.get("contents") or {}
            if contents:
                if "quote" in contents:
                    quoted_content = contents["quote"].get("content", "")
                    msg_content = "[引用]: {}".format(quoted_content)
                elif "forward" in contents:
                    msg_content = "[转发消息]"
                else:
                    msg_content = "[系统消息]"
            else:
                msg_content = "[系统消息]"
        elif msg_type in (50, 11000):
            msg_content = "[语音电话]"
        elif msg_type == 10000:
            msg_content = "[撤回消息]"
        else:
            msg_content = "[其他消息]"

        optimized_messages.append({
            "time": formatted_time,
            "sender": formatted_sender,
            "message": msg_content,
        })

    return optimized_messages


# ============================================================
# 三、统计分析（语音电话 / 常用词 / 趋势 / 频次 / 月占比）
# ============================================================

SPECIAL_MESSAGES = [
    '[photo]', '[voice]', '[video]', '[emoji]', '[引用]',
    '[转发消息]', '[语音电话]', '[撤回消息]', '[其他消息]', '[系统消息]'
]

STOP_WORDS = {
    '的', '了', '是', '在', '我', '你', '他', '她', '它', '我们', '你们', '他们',
    '这', '那', '这个', '那个', '就', '都', '也', '还', '又', '再', '很', '非常',
    '啊', '吧', '呢', '吗', '哦', '嗯', '哈', '嘿', '呀', '啦', '嘛', '呗',
    '不', '没', '别', '不要', '不是', '没有', '好', '对', '行', '可以', '知道',
    '什么', '怎么', '为什么', '哪里', '谁', '哪个', '多少', '几', '时候',
    '去', '来', '到', '从', '向', '往', '给', '为', '把', '被', '让',
    '说', '想', '看', '做', '要', '会', '能', '可以', '应该', '需要', '想要',
    '个', '次', '天', '年', '月', '日', '时', '分', '秒', '点', '些', '种',
    '一个', '两个', '三个', '几个', '很多', '一些', '所有', '全部', '每个',
    '因为', '所以', '但是', '不过', '虽然', '可是', '然后', '接着', '最后',
    '或者', '还是', '如果', '要是', '只要', '只有', '无论', '不管', '即使',
    '感觉', '觉得', '认为', '以为', '知道', '明白', '理解', '清楚', '记得',
    '现在', '以前', '以后', '刚才', '马上', '立刻', '很快', '已经', '还是',
    '真的', '确实', '当然', '肯定', '一定', '必须', '应该', '可能', '也许',
    '这样', '那样', '怎么样', '如何', '什么', '哪里', '谁', '为什么',
    '哈哈', '嘿嘿', '呵呵', '嘻嘻', '哎呀', '嗯嗯', '好的', '行吧', '可以啊',
    '吗', '呢', '吧', '啊', '哦', '呀', '啦', '嘛', '呗', '嗯', '哈', '嘿'
}


def _parse_time(time_str):
    return datetime.strptime(time_str, '%Y-%m-%d %H:%M:%S')


def process_voice_calls(chat_data):
    """语音通话记录表"""
    voice_call_records = []
    for message in chat_data:
        if message['message'] == '[语音电话]':
            datetime_obj = _parse_time(message['time'])
            voice_call_records.append({
                '日期': datetime_obj.date(),
                '时间': datetime_obj.time(),
                '发送者': message['sender'],
                '完整时间': message['time']
            })

    df = pd.DataFrame(voice_call_records)
    if not df.empty:
        df = df.sort_values('完整时间').reset_index(drop=True)
        df.insert(0, '序号', range(1, len(df) + 1))
    return df


def extract_text_messages(chat_data):
    """过滤特殊消息，提取纯文本"""
    text_messages = []
    for message in chat_data:
        msg = message['message']
        if msg not in SPECIAL_MESSAGES and not msg.startswith('[引用]'):
            text_messages.append(msg)
    return text_messages


def process_word_frequency(text_messages):
    """分词并统计 TOP50 常用词"""
    all_text = ' '.join(text_messages)
    words = jieba.lcut(all_text)

    filtered_words = []
    for word in words:
        if (len(word) > 1 and
                word not in STOP_WORDS and
                not word.isdigit() and
                not re.match(r'^[^\w\u4e00-\u9fa5]+$', word)):
            filtered_words.append(word)

    word_counts = Counter(filtered_words)
    top_50 = word_counts.most_common(50)

    df = pd.DataFrame(top_50, columns=['词语', '出现次数'])
    df.insert(0, '排名', range(1, len(df) + 1))
    return df


def process_top_10_days(chat_data):
    """聊天数最高 10 日"""
    dates = [_parse_time(m['time']).date() for m in chat_data]
    df = pd.DataFrame({'date': dates})
    daily_stats = df.groupby('date').size().reset_index()
    daily_stats.columns = ['日期', '消息数量']

    top_10 = daily_stats.sort_values('消息数量', ascending=False).head(10).reset_index(drop=True)
    top_10.insert(0, '排名', range(1, len(top_10) + 1))
    return top_10


def process_hourly_trend(chat_data):
    """24 小时消息分布"""
    dates = [_parse_time(m['time']) for m in chat_data]
    df = pd.DataFrame({'datetime': dates})
    df['hour'] = df['datetime'].dt.hour

    hourly_stats = df.groupby('hour').size().reset_index()
    hourly_stats.columns = ['hour', 'count']

    all_hours = pd.DataFrame({'hour': range(24)})
    hourly_stats = all_hours.merge(hourly_stats, on='hour', how='left').fillna(0)
    hourly_stats.columns = ['小时', '消息数量']
    hourly_stats['消息数量'] = hourly_stats['消息数量'].astype(int)
    return hourly_stats


def process_weekday_trend(chat_data):
    """周内（周一至周日）消息分布"""
    dates = [_parse_time(m['time']) for m in chat_data]
    df = pd.DataFrame({'datetime': dates})
    df['weekday'] = df['datetime'].dt.weekday

    weekly_stats = df.groupby('weekday').size().reset_index()
    weekly_stats.columns = ['weekday', 'count']

    all_weekdays = pd.DataFrame({'weekday': range(7)})
    weekly_stats = all_weekdays.merge(weekly_stats, on='weekday', how='left').fillna(0)
    weekly_stats.columns = ['星期', '消息数量']
    weekly_stats['消息数量'] = weekly_stats['消息数量'].astype(int)

    weekday_names = ['周一', '周二', '周三', '周四', '周五', '周六', '周日']
    weekly_stats['星期'] = weekly_stats['星期'].map(lambda x: weekday_names[x])
    return weekly_stats


def process_monthday_trend(chat_data):
    """月内（1-31 日）消息分布"""
    dates = [_parse_time(m['time']) for m in chat_data]
    df = pd.DataFrame({'datetime': dates})
    df['day'] = df['datetime'].dt.day

    monthly_stats = df.groupby('day').size().reset_index()
    monthly_stats.columns = ['day', 'count']

    all_days = pd.DataFrame({'day': range(1, 32)})
    monthly_stats = all_days.merge(monthly_stats, on='day', how='left').fillna(0)
    monthly_stats.columns = ['日期', '消息数量']
    monthly_stats['消息数量'] = monthly_stats['消息数量'].astype(int)
    return monthly_stats


def process_weekly_data(chat_data):
    """逐周消息数量"""
    dates = [_parse_time(m['time']) for m in chat_data]
    df = pd.DataFrame({'datetime': dates})
    df['year_week'] = df['datetime'].dt.strftime('%Y-%W')
    df['week_start'] = df['datetime'].dt.to_period('W').apply(lambda x: x.start_time.date())
    df['week_end'] = df['datetime'].dt.to_period('W').apply(lambda x: x.end_time.date())

    weekly_stats = df.groupby(['year_week', 'week_start', 'week_end']).size().reset_index()
    weekly_stats.columns = ['年-周', '周开始日期', '周结束日期', '消息数量']
    return weekly_stats.sort_values('年-周').reset_index(drop=True)


def process_monthly_data(chat_data):
    """逐月消息数量"""
    dates = [_parse_time(m['time']) for m in chat_data]
    df = pd.DataFrame({'datetime': dates})
    df['year_month'] = df['datetime'].dt.strftime('%Y-%m')

    monthly_stats = df.groupby('year_month').size().reset_index()
    monthly_stats.columns = ['年月', '消息数量']
    return monthly_stats.sort_values('年月').reset_index(drop=True)


def process_monthly_percentage(chat_data):
    """每月聊天天数占当月总天数的比例"""
    dates = [_parse_time(m['time']).date() for m in chat_data]
    df = pd.DataFrame({'date': dates})
    df['date'] = pd.to_datetime(df['date'])
    df['year_month'] = df['date'].dt.strftime('%Y-%m')

    monthly_chat_days = df.groupby('year_month')['date'].nunique().reset_index()
    monthly_chat_days.columns = ['年月', '聊天天数']
    monthly_chat_days['总天数'] = monthly_chat_days['年月'].apply(lambda x:
        calendar.monthrange(int(x.split('-')[0]), int(x.split('-')[1]))[1])
    monthly_chat_days['占比(%)'] = (monthly_chat_days['聊天天数'] /
                                    monthly_chat_days['总天数']).apply(lambda x: "{:.2%}".format(x))
    return monthly_chat_days


# ------------------------------------------------------------
# 图表生成（每张工作表的数据对应生成图表）
# ------------------------------------------------------------

WORD_CHART_TOP_N = 15  # 常用词图表展示的词语数量


def _add_chart(ws, kind, title, cat_col, val_col, n_rows, anchor, x_title, y_title,
               last_val_col=None, top_n=None, width=18, height=9):
    """为工作表数据生成图表并锚定到 anchor 单元格。

    数据布局约定：第 1 行为表头，第 2 至 n_rows+1 行为数据。
    kind：'bar' 柱状图 / 'hbar' 横向条形图 / 'line' 折线图；
    val_col..last_val_col 为数值列（多列时生成多条数据系列）；
    top_n 指定时仅取前 top_n 行数据作图。
    """
    if n_rows <= 0:
        return
    last_val_col = last_val_col or val_col
    last_row = (n_rows if top_n is None else min(n_rows, top_n)) + 1
    data = Reference(ws, min_col=val_col, max_col=last_val_col, min_row=1, max_row=last_row)
    cats = Reference(ws, min_col=cat_col, min_row=2, max_row=last_row)

    if kind == 'line':
        chart = LineChart()
    else:
        chart = BarChart()
        chart.type = 'bar' if kind == 'hbar' else 'col'
        chart.gapWidth = 60

    chart.title = title
    chart.style = 10
    chart.width = width
    chart.height = height
    chart.add_data(data, titles_from_data=True)
    chart.set_categories(cats)
    chart.x_axis.title = x_title
    chart.y_axis.title = y_title

    if kind == 'line':
        for series in chart.series:
            series.smooth = False
            series.marker = Marker(symbol='circle', size=5)
    else:
        if kind == 'hbar':
            # 反转分类方向，使第一行数据显示在最上方
            chart.x_axis.scaling.orientation = 'maxMin'
            chart.y_axis.crosses = 'max'
        chart.dataLabels = DataLabelList()
        chart.dataLabels.showVal = True

    ws.add_chart(chart, anchor)


def _add_voice_call_chart(ws, df):
    """在语音电话表下方追加"按日期统计通话次数"数据块，并生成对应柱状图"""
    daily = df['日期'].value_counts().sort_index()

    label_row = len(df) + 3  # 主表占 1 行表头 + len(df) 行数据，空 1 行后开始
    header_row = label_row + 1
    ws.cell(row=label_row, column=1, value='按日期统计通话次数')
    ws.cell(row=header_row, column=1, value='日期')
    ws.cell(row=header_row, column=2, value='通话次数')
    for i, (day, count) in enumerate(daily.items()):
        ws.cell(row=header_row + 1 + i, column=1, value=str(day))
        ws.cell(row=header_row + 1 + i, column=2, value=int(count))

    chart = BarChart()
    chart.type = 'col'
    chart.title = '每日语音通话次数'
    chart.style = 10
    chart.gapWidth = 60
    chart.width = 18
    chart.height = 9
    data = Reference(ws, min_col=2, min_row=header_row, max_row=header_row + len(daily))
    cats = Reference(ws, min_col=1, min_row=header_row + 1, max_row=header_row + len(daily))
    chart.add_data(data, titles_from_data=True)
    chart.set_categories(cats)
    chart.x_axis.title = '日期'
    chart.y_axis.title = '通话次数'
    chart.dataLabels = DataLabelList()
    chart.dataLabels.showVal = True
    ws.add_chart(chart, 'G2')


def run_all_stats(chat_data, output_file):
    """执行全量统计分析并写入 Excel，每张工作表附带对应图表"""
    frames = {}

    with pd.ExcelWriter(output_file, engine='openpyxl') as writer:
        voice_call_data = process_voice_calls(chat_data)
        if not voice_call_data.empty:
            voice_call_data.to_excel(writer, sheet_name='语音电话表', index=False)
            frames['语音电话表'] = voice_call_data
            _add_voice_call_chart(writer.sheets['语音电话表'], voice_call_data)
            print("  - 语音电话统计：{} 条记录（含每日通话次数图表）".format(len(voice_call_data)))

        word_stats = process_word_frequency(extract_text_messages(chat_data))
        word_stats.to_excel(writer, sheet_name='常用词表', index=False)
        frames['常用词表'] = word_stats
        word_chart_n = min(len(word_stats), WORD_CHART_TOP_N)
        _add_chart(writer.sheets['常用词表'], 'hbar', '常用词语 TOP{}'.format(word_chart_n),
                   cat_col=2, val_col=3, n_rows=len(word_stats), anchor='E2',
                   x_title='词语', y_title='出现次数', top_n=WORD_CHART_TOP_N, height=12)
        print("  - 常用词语统计：TOP{}（含词频图表）".format(len(word_stats)))

        top_10_stats = process_top_10_days(chat_data)
        top_10_stats.to_excel(writer, sheet_name='聊天数最高10日表', index=False)
        frames['聊天数最高10日表'] = top_10_stats
        _add_chart(writer.sheets['聊天数最高10日表'], 'bar', '聊天数最高10日',
                   cat_col=2, val_col=3, n_rows=len(top_10_stats), anchor='E2',
                   x_title='日期', y_title='消息数量')
        print("  - 聊天数最高10日统计完成（含图表）")

        hourly_stats = process_hourly_trend(chat_data)
        hourly_stats.to_excel(writer, sheet_name='24小时趋势表', index=False)
        _add_chart(writer.sheets['24小时趋势表'], 'line', '24小时消息分布',
                   cat_col=1, val_col=2, n_rows=len(hourly_stats), anchor='D2',
                   x_title='小时', y_title='消息数量', width=20)

        weekday_stats = process_weekday_trend(chat_data)
        weekday_stats.to_excel(writer, sheet_name='周趋势表', index=False)
        _add_chart(writer.sheets['周趋势表'], 'bar', '一周消息分布（周一至周日）',
                   cat_col=1, val_col=2, n_rows=len(weekday_stats), anchor='D2',
                   x_title='星期', y_title='消息数量')

        monthday_stats = process_monthday_trend(chat_data)
        monthday_stats.to_excel(writer, sheet_name='月趋势表', index=False)
        _add_chart(writer.sheets['月趋势表'], 'line', '月内每日消息分布（1-31日）',
                   cat_col=1, val_col=2, n_rows=len(monthday_stats), anchor='D2',
                   x_title='日期', y_title='消息数量', width=20)
        print("  - 24小时/周/月趋势统计完成（含图表）")

        weekly_data = process_weekly_data(chat_data)
        weekly_data.to_excel(writer, sheet_name='周消息数量表', index=False)
        _add_chart(writer.sheets['周消息数量表'], 'line', '每周消息数量趋势',
                   cat_col=1, val_col=4, n_rows=len(weekly_data), anchor='F2',
                   x_title='周', y_title='消息数量', width=20)

        monthly_data = process_monthly_data(chat_data)
        monthly_data.to_excel(writer, sheet_name='月消息数量表', index=False)
        _add_chart(writer.sheets['月消息数量表'], 'bar', '每月消息数量',
                   cat_col=1, val_col=2, n_rows=len(monthly_data), anchor='D2',
                   x_title='月份', y_title='消息数量')
        print("  - 周/月消息频次统计完成（含图表）")

        monthly_percentage = process_monthly_percentage(chat_data)
        monthly_percentage.to_excel(writer, sheet_name='消息月占比', index=False)
        _add_chart(writer.sheets['消息月占比'], 'bar', '每月聊天天数与当月总天数',
                   cat_col=1, val_col=2, last_val_col=3, n_rows=len(monthly_percentage),
                   anchor='F2', x_title='月份', y_title='天数')
        print("  - 消息月占比统计完成（含图表）")

        # 自动调整列宽
        for sheet_name in writer.sheets:
            worksheet = writer.sheets[sheet_name]
            for column_cells in worksheet.columns:
                length = max(len(str(cell.value)) for cell in column_cells)
                worksheet.column_dimensions[column_cells[0].column_letter].width = length + 2

    return frames


# ============================================================
# 四、心理测评任务包（供主智能体派发子智能体执行）
# ============================================================

PSYCH_SYSTEM_PROMPT = """你是一位专业的心理测评专家，擅长基于日常聊天记录分析个体的情绪状态与心理特征。请严格遵守以下原则：
1. 仅基于聊天记录中的客观内容进行分析，不做无依据的推测；
2. 测评结果仅供自我了解、情绪关怀与关系参考，不构成临床心理诊断；
3. 语言专业、温和、具有建设性，避免评判性表述。"""

PSYCH_USER_TEMPLATE = """请对以下一周的聊天记录进行心理测评，并按模板输出结构化的 Markdown 测评报告。

【测评报告模板】
## 一、情绪状态分析
（积极 / 中性 / 消极情绪的分布与大致占比，本周情绪整体基调）
## 二、情绪变化趋势
（本周内情绪的起伏、转折点及其触发原因）
## 三、关键心理事件
（引发明显情绪波动的重要对话，可引用原文关键句）
## 四、心理特征观察
（压力水平、焦虑信号、情绪稳定性、社交互动模式等维度）
## 五、综合测评结论
（用 3-5 句话总结本周整体心理状态）
## 六、关怀建议
（给出 2-3 条具体、可操作的建议）

聊天记录时间范围：{start_time} 到 {end_time}
聊天记录（共 {count} 条）：
{chat_text}"""


def chat_to_dataframe(chat_data):
    df = pd.DataFrame(chat_data)
    df['datetime'] = pd.to_datetime(df['time'])
    df['date'] = df['datetime'].dt.date
    df['week'] = df['datetime'].dt.to_period('W')
    return df


def group_by_week(df):
    """按周分组，返回每周的聊天记录"""
    weekly_chats = []
    for week, group in df.groupby('week'):
        week_chats = group.sort_values('datetime')
        weekly_chats.append({
            'week': week,
            'start_time': week_chats['datetime'].min().strftime('%Y-%m-%d %H:%M:%S'),
            'end_time': week_chats['datetime'].max().strftime('%Y-%m-%d %H:%M:%S'),
            'messages': week_chats.to_dict('records'),
            'count': len(week_chats)
        })
    weekly_chats.sort(key=lambda x: x['week'].start_time)
    return weekly_chats


def format_chat_for_ai(messages):
    formatted = []
    for msg in messages:
        formatted.append("[{}] {}: {}".format(msg['time'], msg['sender'], msg['message']))
    return '\n'.join(formatted)


def build_assessment_tasks(chat_data, source):
    """将聊天记录按周分批，生成心理测评任务包（由外部子智能体执行）"""
    df = chat_to_dataframe(chat_data)
    weekly_chats = group_by_week(df)

    tasks = []
    for i, week in enumerate(weekly_chats):
        chat_text = format_chat_for_ai(week['messages'])
        tasks.append({
            "task_type": "psychological_assessment",
            "task_id": "batch_{:03d}".format(i + 1),
            "batch": i + 1,
            "week": str(week['week']),
            "time_range": {"start": week['start_time'], "end": week['end_time']},
            "message_count": week['count'],
            "result_file_suggestion": "assessment_results/batch_{:03d}.md".format(i + 1),
            "sub_agent": {
                "role": "心理测评子智能体",
                "system_prompt": PSYCH_SYSTEM_PROMPT,
                "user_prompt": PSYCH_USER_TEMPLATE.format(
                    start_time=week['start_time'],
                    end_time=week['end_time'],
                    count=week['count'],
                    chat_text=chat_text,
                ),
            },
        })

    return {
        "generated_at": datetime.now().strftime("%Y-%m-%d %H:%M:%S"),
        "source": source,
        "total_tasks": len(tasks),
        "instructions_for_main_agent": (
            "你是主智能体，负责监测与调度：\n"
            "1. 调用统计脚本 python main.py <URL>，通过 run_summary.json 监测各阶段状态；\n"
            "2. 统计完成后读取本文件，将 tasks 数组中的任务逐个派发给心理测评子智能体；\n"
            "3. 收集子智能体产出的测评报告，按 task_id / 时间顺序汇总为最终心理测评总报告。"
        ),
        "instructions_for_sub_agent": (
            "你是心理测评子智能体，收到一个任务包后：\n"
            "1. 将 task.sub_agent.system_prompt 设为你的系统提示词；\n"
            "2. 将 task.sub_agent.user_prompt 作为用户消息（其中已包含完整聊天记录）；\n"
            "3. 按模板输出结构化心理测评报告（Markdown），保存到 task.result_file_suggestion "
            "指示的位置（相对于统计输出目录），或交回主智能体汇总。"
        ),
        "tasks": tasks,
    }


# ============================================================
# 五、运行摘要
# ============================================================

def build_stats_summary(chat_data, frames, task_count):
    if not chat_data:
        return {}
    times = [m['time'] for m in chat_data if m.get('time')]
    senders = Counter(m['sender'] for m in chat_data)

    voice_df = frames.get('语音电话表')
    word_df = frames.get('常用词表')
    days_df = frames.get('聊天数最高10日表')

    return {
        "message_count": len(chat_data),
        "time_range": {"start": min(times), "end": max(times)},
        "sender_distribution": dict(senders),
        "voice_call_count": 0 if voice_df is None else len(voice_df),
        "top_10_words": [
            {"排名": int(r['排名']), "词语": r['词语'], "出现次数": int(r['出现次数'])}
            for r in (word_df.head(10).to_dict('records') if word_df is not None else [])
        ],
        "top_5_chat_days": [
            {"排名": int(r['排名']), "日期": str(r['日期']), "消息数量": int(r['消息数量'])}
            for r in (days_df.head(5).to_dict('records') if days_df is not None else [])
        ],
        "assessment_task_count": task_count,
    }


# ============================================================
# 六、主流程：输入一个 URL，返回所有结果
# ============================================================

def main():
    parser = argparse.ArgumentParser(
        description="微信聊天记录分析工具（单文件版）：输入一个 chatlog API URL，"
                    "自动完成下载、格式化、全量统计，并生成心理测评任务包。",
        formatter_class=argparse.RawDescriptionHelpFormatter,
        epilog='示例：\n'
               '  python main.py "http://127.0.0.1:5030/api/v1/chatlog?limit=100000"\n'
               '  python main.py chat.json -o output\n'
               '注意：URL 含 & 等特殊字符时请用引号包裹。'
    )
    parser.add_argument('source', help='chatlog API 地址（或本地聊天记录 JSON 文件路径）')
    parser.add_argument('-o', '--output', default='output', help='输出目录（默认：output）')
    parser.add_argument('--timeout', type=int, default=60, help='URL 下载超时秒数（默认：60）')
    args = parser.parse_args()

    output_dir = args.output
    os.makedirs(output_dir, exist_ok=True)
    outputs = {
        "formatted_chat": os.path.join(output_dir, "chat_format.json"),
        "statistics": os.path.join(output_dir, "chat_analysis.xlsx"),
        "assessment_tasks": os.path.join(output_dir, "assessment_tasks.json"),
        "run_summary": os.path.join(output_dir, "run_summary.json"),
    }

    monitor = RunMonitor(args.source, output_dir)

    print("=" * 60)
    print("微信聊天记录分析工具（单文件版）")
    print("=" * 60)

    try:
        # 阶段 1：数据下载
        monitor.report("数据下载", "running")
        raw_messages = load_chat_messages(args.source, args.timeout)
        monitor.report("数据下载", "done", "共获取 {} 条原始消息".format(len(raw_messages)))

        # 阶段 2：格式化
        monitor.report("格式化聊天记录", "running")
        chat_data = format_chat_messages(raw_messages)
        with open(outputs["formatted_chat"], "w", encoding="utf-8") as f:
            json.dump(chat_data, f, ensure_ascii=False, indent=2)
        monitor.report("格式化聊天记录", "done",
                       "共 {} 条，已保存 {}".format(len(chat_data), outputs["formatted_chat"]))

        # 阶段 3：全量统计
        monitor.report("全量统计分析", "running")
        frames = run_all_stats(chat_data, outputs["statistics"])
        monitor.report("全量统计分析", "done",
                       "已生成 {}（每张工作表附带对应图表）".format(outputs["statistics"]))

        # 阶段 4：心理测评任务包
        monitor.report("生成心理测评任务包", "running")
        task_doc = build_assessment_tasks(chat_data, args.source)
        with open(outputs["assessment_tasks"], "w", encoding="utf-8") as f:
            json.dump(task_doc, f, ensure_ascii=False, indent=2)
        monitor.report("生成心理测评任务包", "done",
                       "共 {} 个周批次任务，已保存 {}".format(
                           task_doc["total_tasks"], outputs["assessment_tasks"]))

        # 运行摘要（供主智能体监测）
        stats_summary = build_stats_summary(chat_data, frames, task_doc["total_tasks"])
        summary = monitor.summary("success", outputs, stats_summary)
        with open(outputs["run_summary"], "w", encoding="utf-8") as f:
            json.dump(summary, f, ensure_ascii=False, indent=2)

        print()
        print("=" * 60)
        print("全部完成！输出文件：")
        for path in outputs.values():
            print("  - {}".format(path))
        print("=" * 60)
        print()
        print("消息总数：{message_count}，时间范围：{start} ~ {end}".format(
            message_count=stats_summary.get("message_count", 0),
            start=stats_summary.get("time_range", {}).get("start", ""),
            end=stats_summary.get("time_range", {}).get("end", "")))
        print()
        print("下一步（主智能体）：派发心理测评子智能体，读取 assessment_tasks.json，")
        print("逐个执行 tasks 中的测评任务（包内含完整提示词与聊天记录），")
        print("并按 task.result_file_suggestion 汇总测评报告。")
        return 0

    except Exception as e:
        monitor.fail(str(e))
        summary = monitor.summary("failed", outputs)
        try:
            with open(outputs["run_summary"], "w", encoding="utf-8") as f:
                json.dump(summary, f, ensure_ascii=False, indent=2)
        except OSError:
            pass
        print("\n执行失败：{}".format(e), file=sys.stderr)
        print("失败详情已写入 {}".format(outputs["run_summary"]), file=sys.stderr)
        return 1


if __name__ == "__main__":
    sys.exit(main())
